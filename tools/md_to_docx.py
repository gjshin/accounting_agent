"""검토 의견서·조서 메모(마크다운)를 Word 문서로 바꾼다.

사용법:
    python3 tools/md_to_docx.py <입력.md> <출력.docx> [--header "머리글 문구"] [--subtitle "부제"] [--date YYYY-MM-DD]

- 첫 줄의 '# 제목'을 표지 제목으로 쓰고, 표지에 목차(1·2단계 제목)를 글자로 넣는다.
  (Word의 자동 목차 필드는 뷰어에 따라 비어 보여서 쓰지 않는다.)
- 마크다운 각주([^이름])는 Word 각주로 바뀐다.
- 표에는 테두리와 머리행 음영을, 인용 블록(기준서·법령 원문)에는 들여쓰기와 음영을 넣는다.
- 회사 양식이 있으면 templates/reference.docx 에 두면 그 양식(글꼴·머리글·여백)을 쓴다.
  없으면 맑은 고딕 기본 양식을 만들어 쓴다.
- pandoc 이 없으면 pypandoc_binary 를, python-docx 가 없으면 python-docx 를 설치한다.
"""
import argparse
import datetime as dt
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
USER_TEMPLATE = ROOT / "templates" / "reference.docx"
FONT = "맑은 고딕"


def ensure(module: str, package: str):
    try:
        return __import__(module)
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", package], check=True)
        return __import__(module)


pypandoc = ensure("pypandoc", "pypandoc_binary")
docx = ensure("docx", "python-docx")
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402


# Word는 XML 하위 요소의 순서를 엄격히 따진다. 음영(shd) 뒤에 와야 하는 요소들.
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                 "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                 "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                 "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                  "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                       "w:tblDescription", "w:tblPrChange")


def set_font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def shade(parent, fill, successors):
    for old in parent.findall(qn("w:shd")):
        parent.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    parent.insert_element_before(shd, *successors)


def add_page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build_reference(path: Path, header_text: str) -> None:
    subprocess.run(
        [pypandoc.get_pandoc_path(), "-o", str(path), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    d = docx.Document(str(path))
    styles = d.styles
    for name, size in (("Normal", 10.5), ("Body Text", 10.5), ("First Paragraph", 10.5),
                       ("Compact", 10.5), ("Footnote Text", 8.5)):
        if name in [s.name for s in styles]:
            set_font(styles[name], size)
    for name, size in (("Title", 22), ("Subtitle", 13), ("Date", 11),
                       ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        if name in [s.name for s in styles]:
            set_font(styles[name], size, bold=name.startswith("Heading") or name == "Title", color="1F3864")
    if "Block Text" in [s.name for s in styles]:
        bt = styles["Block Text"]
        set_font(bt, 10)
        bt.paragraph_format.left_indent = Pt(14)
        bt.paragraph_format.right_indent = Pt(6)
        shade(bt.element.get_or_add_pPr(), "F2F2F2", PPR_AFTER_SHD)
    section = d.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = ""
    run = hp.add_run(header_text)
    run.font.size = Pt(8.5)
    fp = section.footer.paragraphs[0]
    fp.alignment = 1
    add_page_number(fp)
    d.save(str(path))


def style_tables(d) -> None:
    for table in d.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        for old in tbl_pr.findall(qn("w:tblBorders")):
            tbl_pr.remove(old)
        tbl_pr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)
        for cell in table.rows[0].cells:
            shade(cell._tc.get_or_add_tcPr(), "D9E2F3", TCPR_AFTER_SHD)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True


def page_breaks(d) -> None:
    first_h1 = True
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            if first_h1 or p.text.strip().startswith("부록"):
                p.paragraph_format.page_break_before = True
            first_h1 = False


def static_toc(body: str) -> str:
    lines = ['::: {custom-style="TOC Heading"}', "목차", ":::", ""]
    in_code = False
    for line in body.splitlines():
        if line.startswith("```"):
            in_code = not in_code
        m = None if in_code else re.match(r"^(#{1,2})\s+(.+)$", line)
        if m:
            style = "TOC 1" if len(m.group(1)) == 1 else "TOC 2"
            text = re.sub(r"^(\d+)\.", r"\1\\.", m.group(2).strip())  # '1. 요약'이 번호 목록으로 바뀌지 않게
            lines += [f'::: {{custom-style="{style}"}}', text, ":::", ""]
    return "\n".join(lines) + "\n"


def convert(src: Path, dst: Path, header: str, subtitle: str, date: str) -> None:
    text = src.read_text(encoding="utf-8")
    m = re.match(r"\s*#\s+(.+)\n", text)
    title = m.group(1).strip() if m else src.stem
    body = text[m.end():] if m else text
    # 본문 '## 1. 요약' 같은 2단계 제목을 Word의 1단계 제목으로 올린다.
    body = re.sub(r"^(#{2,6})\s", lambda x: "#" * (len(x.group(1)) - 1) + " ", body, flags=re.MULTILINE)
    body = static_toc(body) + body
    with tempfile.TemporaryDirectory() as tmp:
        ref = USER_TEMPLATE
        if not ref.exists():
            ref = Path(tmp) / "reference.docx"
            build_reference(ref, header)
        md = Path(tmp) / "in.md"
        md.write_text(body, encoding="utf-8")
        pypandoc.convert_file(
            str(md), "docx", format="markdown+footnotes+pipe_tables+hard_line_breaks+fenced_divs", outputfile=str(dst),
            extra_args=[
                f"--reference-doc={ref}",
                "-M", f"title={title}", "-M", f"subtitle={subtitle}", "-M", f"date={date}",
                "-M", "lang=ko-KR",
            ],
        )
    d = docx.Document(str(dst))
    style_tables(d)
    page_breaks(d)
    names = [st.name for st in d.styles]
    if "TOC 2" in names:
        d.styles["TOC 2"].paragraph_format.left_indent = Pt(14)
    d.save(str(dst))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--header", default="대외비 | 회계위키·법령 근거 검토")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--date", default=dt.date.today().isoformat())
    a = ap.parse_args()
    convert(Path(a.src), Path(a.dst), a.header, a.subtitle, a.date)
    print(a.dst)


if __name__ == "__main__":
    main()
